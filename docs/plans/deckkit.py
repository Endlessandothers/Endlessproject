"""Shared deck furniture for the Endless phase decks.

Extracted when the Phase 3 deck needed the same helpers the Phase 2 one had.
Two copies of a layout function is two layouts the moment one is edited, and
these decks are meant to read as one set.

Style is Phase 1's, deliberately not reinvented: 16:9, Segoe UI, slate ink on
white, one orange rule under each title.
"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

ICONS = Path(__file__).resolve().parent / "icons"

INK = RGBColor(0x0F, 0x17, 0x2A)
ORANGE = RGBColor(0xED, 0x71, 0x00)
MUTED = RGBColor(0x94, 0xA3, 0xB8)
BODY = RGBColor(0x47, 0x55, 0x69)
DIM = RGBColor(0x64, 0x74, 0x8B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

CALLOUT_BG = RGBColor(0xFF, 0xF7, 0xED)
CALLOUT_TX = RGBColor(0x7C, 0x2D, 0x12)
DONE_BG = RGBColor(0xF0, 0xFD, 0xF4)
DONE_TX = RGBColor(0x14, 0x53, 0x2D)
WARN_BG = RGBColor(0xFE, 0xF2, 0xF2)
WARN_TX = RGBColor(0x7F, 0x1D, 0x1D)

FONT = "Segoe UI"

# The board's colour code, so a deck teaches the same legend the Miro frame uses.
ZONE = {
    "new":      (RGBColor(0xDC, 0xFC, 0xE7), RGBColor(0x14, 0x53, 0x2D), RGBColor(0x16, 0xA3, 0x4A)),
    "existing": (RGBColor(0xFE, 0xF3, 0xC7), RGBColor(0x78, 0x35, 0x0F), RGBColor(0xF5, 0x9E, 0x0B)),
    "sandbox":  (RGBColor(0xFE, 0xE2, 0xE2), RGBColor(0x7F, 0x1D, 0x1D), RGBColor(0xDC, 0x26, 0x26)),
    "data":     (RGBColor(0xDB, 0xEA, 0xFE), RGBColor(0x1E, 0x3A, 0x8A), RGBColor(0x3B, 0x82, 0xF6)),
    "public":   (RGBColor(0xED, 0xE9, 0xFE), RGBColor(0x4C, 0x1D, 0x95), RGBColor(0x8C, 0x4F, 0xFF)),
}


def icon(name):
    """Official AWS service icons, from docs/plans/icons."""
    return str(ICONS / f"{name}.png")


def deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    return prs


def slide(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = s.background.fill
    bg.solid()
    bg.fore_color.rgb = WHITE
    return s


def text(s, x, y, w, h, content, size, colour, bold=False, spacing=1.15):
    box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(content.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = spacing
        r = p.add_run()
        r.text = line
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.name = FONT
        r.font.color.rgb = colour
    return box


def heading(s, eyebrow, title):
    text(s, 0.85, 0.55, 11.6, 0.3, eyebrow, 11, MUTED, bold=True)
    text(s, 0.85, 0.92, 11.6, 0.7, title, 31, INK, bold=True)
    rule = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.85), Inches(1.72), Inches(1.1), Inches(0.04))
    rule.fill.solid()
    rule.fill.fore_color.rgb = ORANGE
    rule.line.fill.background()
    rule.shadow.inherit = False


def callout(s, x, y, w, h, body, bg=CALLOUT_BG, tx=CALLOUT_TX, size=13.5):
    box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    box.fill.solid()
    box.fill.fore_color.rgb = bg
    box.line.fill.background()
    box.shadow.inherit = False
    box.adjustments[0] = 0.04
    tb = text(s, x + 0.3, y + 0.2, w - 0.6, h - 0.4, body, size, tx, bold=True)
    tb.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    return box


def table(s, x, y, w, rows, widths, sizes=11, header=True):
    shape = s.shapes.add_table(len(rows), len(widths), Inches(x), Inches(y), Inches(w),
                               Inches(0.32 * len(rows)))
    t = shape.table
    t.first_row = header
    for i, cw in enumerate(widths):
        t.columns[i].width = Inches(cw)
    for r, row in enumerate(rows):
        t.rows[r].height = Inches(0.32)
        for c, val in enumerate(row):
            cell = t.cell(r, c)
            cell.text = str(val)
            cell.margin_left = Inches(0.1)
            cell.margin_right = Inches(0.08)
            cell.margin_top = Inches(0.03)
            cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = INK if (header and r == 0) else WHITE
            p = cell.text_frame.paragraphs[0]
            p.line_spacing = 1.0
            for run in p.runs:
                run.font.size = Pt(sizes)
                run.font.name = FONT
                run.font.bold = header and r == 0
                run.font.color.rgb = WHITE if (header and r == 0) else BODY
    return t


def legend_card(s, x, y, w, icon_name, title_, sub, kind):
    """One card in the board's own visual language: icon badge, name, role."""
    bg, tx, line = ZONE[kind]
    card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(0.82))
    card.fill.solid()
    card.fill.fore_color.rgb = bg
    card.line.color.rgb = line
    card.line.width = Pt(1)
    card.shadow.inherit = False
    card.adjustments[0] = 0.08
    s.shapes.add_picture(icon(icon_name), Inches(x + 0.14), Inches(y + 0.18), Inches(0.46), Inches(0.46))
    text(s, x + 0.76, y + 0.11, w - 0.9, 0.28, title_, 13, tx, bold=True)
    text(s, x + 0.76, y + 0.40, w - 0.9, 0.3, sub, 10.5, tx)


def title_slide(prs, eyebrow, title, blurb, footer, icons):
    s = slide(prs)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = INK
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.28), Inches(7.5))
    bar.fill.solid()
    bar.fill.fore_color.rgb = ORANGE
    bar.line.fill.background()
    bar.shadow.inherit = False

    text(s, 1.1, 2.05, 11.0, 0.4, eyebrow, 13, ORANGE, bold=True)
    text(s, 1.1, 2.5, 11.2, 1.2, title, 46, WHITE, bold=True)
    text(s, 1.1, 3.8, 10.4, 0.9, blurb, 17, MUTED)
    for i, name in enumerate(icons):
        s.shapes.add_picture(icon(name), Inches(1.1 + i * 0.72), Inches(5.1), Inches(0.55), Inches(0.55))
    text(s, 1.1, 6.4, 11.5, 0.3, footer, 11.5, DIM)
    return s


def numbered(s, y, num, title_, body_):
    chip = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.85), Inches(y), Inches(0.4), Inches(0.4))
    chip.fill.solid()
    chip.fill.fore_color.rgb = ORANGE
    chip.line.fill.background()
    chip.shadow.inherit = False
    tf = chip.text_frame
    tf.text = str(num)
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    r = tf.paragraphs[0].runs[0]
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.name = FONT
    r.font.color.rgb = WHITE
    text(s, 1.45, y - 0.03, 3.4, 0.4, title_, 15, INK, bold=True)
    text(s, 4.95, y - 0.03, 7.45, 0.9, body_, 13, BODY, spacing=1.25)


def verify(path):
    """Report overflow and overlap rather than opening the file to look."""
    from pptx.util import Emu
    prs = Presentation(path)
    SW, SH = 13.333, 7.5
    problems = []
    for i, s in enumerate(prs.slides, 1):
        boxes = []
        for sh in s.shapes:
            L, T = Emu(sh.left).inches, Emu(sh.top).inches
            W, H = Emu(sh.width).inches, Emu(sh.height).inches
            if T + H > SH + 0.02 or L + W > SW + 0.02:
                problems.append(f"slide {i}: {sh.shape_type} runs off the page")
            if sh.has_table:
                tb = T + sum(Emu(r.height).inches for r in sh.table.rows)
                if tb > SH:
                    problems.append(f"slide {i}: table runs off the bottom at {tb:.2f}")
            if sh.has_text_frame and sh.text_frame.text.strip():
                boxes.append((L, T, W, H, sh.text_frame.text[:24].replace("\n", "|")))
        for a in range(len(boxes)):
            for b in range(a + 1, len(boxes)):
                x1, y1, w1, h1, n1 = boxes[a]
                x2, y2, w2, h2, n2 = boxes[b]
                ox = min(x1 + w1, x2 + w2) - max(x1, x2)
                oy = min(y1 + h1, y2 + h2) - max(y1, y2)
                if ox > 0.25 and oy > 0.15 and not (abs(x1 - x2) < 0.35 and abs(y1 - y2) < 0.35):
                    problems.append(f"slide {i}: {n1!r} overlaps {n2!r}")
    return problems
